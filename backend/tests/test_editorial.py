from io import BytesIO
from docx import Document
from app.documents import render_resume, read_document

def test_profile_links_and_subtitle(client):
    p = client.post('/api/v1/profiles', json={'full_name':'Anna Test', 'subtitle':'Senior Product Engineer'}).json()
    pid = p['id']
    assert p['subtitle'] == 'Senior Product Engineer'
    assert p['links'] == []
    bad = client.post(f'/api/v1/profiles/{pid}/links',json={'label':'Bad','url':'javascript:alert(1)'})
    assert bad.status_code == 422
    a = client.post(f'/api/v1/profiles/{pid}/links',json={'label':'GitHub','url':'https://github.com/anna','position':0})
    assert a.status_code == 201, a.text
    lid = a.json()['id']
    upd = client.patch(f'/api/v1/profiles/{pid}/links/{lid}',json={'label':'Kod','url':'https://github.com/anna','position':1})
    assert upd.status_code == 200
    assert client.get(f'/api/v1/profiles/{pid}').json()['links'][0]['label']=='Kod'
    cv = client.post('/api/v1/resumes',json={'profile_id':pid,'content':{**client.get(f'/api/v1/profiles/{pid}').json(),'template':'editorial'}})
    assert cv.status_code==201,cv.text
    raw=client.get(f"/api/v1/resumes/{cv.json()['id']}/export?format=docx").content
    txt=read_document(raw,'resume.docx')
    assert 'Senior Product Engineer' in txt and 'https://github.com/anna' in txt
    assert client.delete(f'/api/v1/profiles/{pid}/links/{lid}').status_code==204
    assert client.get(f'/api/v1/profiles/{pid}').json()['links']==[]

def test_editorial_styles():
    doc=Document(BytesIO(render_resume({'full_name':'Anna Test','subtitle':'Product Engineer','links':[{'label':'Portfolio','url':'https://example.org','position':0}],'experiences':[{'company':'Example','title':'Engineer','start_date':'2020-01-01'}], 'template':'editorial'})))
    assert doc.paragraphs[0].alignment is not None
    assert 'Product Engineer' in [p.text for p in doc.paragraphs]
    assert 'BBBBBB' in doc.styles['Heading 1'].element.xml
