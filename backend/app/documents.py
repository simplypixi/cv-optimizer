from io import BytesIO
from pathlib import Path
import fitz
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


def read_document(data: bytes, filename: str):
    ext=Path(filename).suffix.lower()
    if ext=='.pdf':
        with fitz.open(stream=data,filetype='pdf') as pdf:
            if len(pdf)>30:raise ValueError('Maksymalnie 30 stron')
            text='\n'.join(page.get_text(sort=True) for page in pdf)
    elif ext=='.docx':
        d=Document(BytesIO(data))
        chunks=[p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:chunks.extend(c.text for c in row.cells)
        for section in d.sections:
            chunks.extend(p.text for p in section.header.paragraphs)
            chunks.extend(p.text for p in section.footer.paragraphs)
        text='\n'.join(chunks)
    else:raise ValueError('Obsługiwane formaty: PDF i DOCX')
    if not text.strip():raise ValueError('Nie znaleziono tekstu — skany i obrazy wymagają OCR (poza MVP).')
    return text[:100000]

def _same_company(name):
    return ' '.join(name.casefold().strip().split())


def _experience_groups(experiences, group=False):
    """Stable employer grouping; never merges original job records."""
    roles = sorted(experiences, key=lambda x: str(x.get('start_date', '')), reverse=True)
    if not group:
        return [{'company': x.get('company', ''), 'roles': [x]} for x in roles]
    output = []
    lookup = {}
    for role in roles:
        key = _same_company(role.get('company', ''))
        if key and key in lookup:
            lookup[key]['roles'].append(role)
        else:
            entry = {'company': role.get('company', ''), 'roles': [role]}
            output.append(entry)
            if key:
                lookup[key] = entry
    return output


def render_resume(content:dict):
    """Return plain-text, single-column, selectable DOCX. Grouping is display-only."""
    from docx.shared import RGBColor
    template = content.get('template', 'classic')
    if template not in ('classic', 'editorial', 'modern', 'compact'):
        raise ValueError('Nieznany szablon CV')
    specs = {
        'classic': dict(font='Carlito', size=10.5, margin=2.0, space=5, heading=12, accent=(32, 44, 60)),
        'editorial': dict(font='Liberation Sans', size=10.3, margin=2.15, space=6, heading=11.3, accent=(45, 45, 45)),
        'compact': dict(font='Liberation Sans', size=9.5, margin=1.5, space=3, heading=10.5, accent=(38, 38, 38)),
    }
    if template == 'modern': template = 'editorial'  # backwards compatibility
    cfg=specs[template]
    doc=Document()
    section=doc.sections[0]
    section.top_margin=Cm(cfg['margin']); section.bottom_margin=Cm(cfg['margin'])
    section.left_margin=Cm(cfg['margin']); section.right_margin=Cm(cfg['margin'])
    normal=doc.styles['Normal']; normal.font.name=cfg['font']; normal.font.size=Pt(cfg['size'])
    normal.paragraph_format.space_after=Pt(cfg['space'])
    for key in ('Title','Heading 1','Heading 2'):
        style=doc.styles[key]; style.font.name=cfg['font']; style.font.color.rgb=RGBColor(*cfg['accent'])
    doc.styles['Title'].font.size=Pt(19 if template != 'compact' else 16)
    doc.styles['Heading 1'].font.size=Pt(cfg['heading']); doc.styles['Heading 2'].font.size=Pt(cfg['size']+0.5)
    title=doc.add_paragraph(style='Title'); title.add_run(content['full_name']).bold=True
    title.alignment=WD_ALIGN_PARAGRAPH.CENTER if template=='classic' else WD_ALIGN_PARAGRAPH.LEFT
    contact='  |  '.join(str(content.get(x) or '') for x in ('email','phone','location') if content.get(x))
    if contact:
        p=doc.add_paragraph(contact)
        p.alignment=title.alignment
    subtitle = content.get('subtitle') or content.get('headline')
    if subtitle:
        p=doc.add_paragraph();p.add_run(subtitle).bold=True
        p.alignment=title.alignment
    if content.get('links'):
        for link in sorted(content['links'], key=lambda v: v.get('position', 0)):
            label = str(link.get('label') or '').strip()
            url = str(link.get('url') or '').strip()
            if url:
                para = doc.add_paragraph(f'{label}: {url}' if label else url)
                para.alignment=title.alignment
    if template == 'editorial':
        from docx.enum.text import WD_BREAK
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        title.paragraph_format.space_after = Pt(3)
        for heading_style in ['Heading 1']:
            style = doc.styles[heading_style]
            style.font.bold = True
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(4)
            pPr = style.element.get_or_add_pPr()
            borders = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            for key, value in [('val','single'),('sz','4'),('space','5'),('color','BBBBBB')]:
                bottom.set(qn('w:'+key), value)
            borders.append(bottom); pPr.append(borders)
    if content.get('summary'):
        doc.add_heading('Professional Summary',1);doc.add_paragraph(content['summary'])
    if content.get('experiences'):
        doc.add_heading('Professional Experience',1)
        for group in _experience_groups(content['experiences'],content.get('career_path',False)):
            grouped=len(group['roles'])>1
            if grouped:
                doc.add_heading(group['company'],2)
            for role in group['roles']:
                heading=role['title'] if grouped else f"{role['title']} — {role['company']}"
                para=doc.add_paragraph()
                para.paragraph_format.keep_with_next=True
                para.add_run(heading).bold=True
                para.add_run(f"  |  {role['start_date']} – {role.get('end_date') or 'Present'}")
                if role.get('description'):
                    doc.add_paragraph(role['description'])
    if content.get('projects'):
        doc.add_heading('Projects',1)
        for item in content['projects']:
            p=doc.add_paragraph();p.add_run(item['name']).bold=True
            if item.get('description'):doc.add_paragraph(item['description'])
            if item.get('url'):doc.add_paragraph(item['url'])
    if content.get('education'):
        doc.add_heading('Education',1)
        for item in content['education']:
            p=doc.add_paragraph();p.add_run(item['institution']).bold=True
            if item.get('degree'):doc.add_paragraph(item['degree'])
    if content.get('skills'):
        doc.add_heading('Skills',1)
        doc.add_paragraph(' · '.join(item['name'] for item in content['skills']))
    out=BytesIO();doc.save(out);return out.getvalue()

def lint(text):
    lines=[x.strip() for x in text.splitlines() if x.strip()]
    warnings=[]
    if len(text.strip())<100:warnings.append('Very little text extracted; verify completeness and reading order.')
    if '@' not in text:warnings.append('No email address detected.')
    if not any(x.lower() in text.lower() for x in ['doświadczenie','experience','employment']):warnings.append('No standard experience section detected.')
    return {'characters':len(text),'lines':len(lines),'warnings':warnings,'note':'Text readability check, not an emulation of any employer ATS.'}
