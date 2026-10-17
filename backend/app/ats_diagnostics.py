"""On-demand, evidence-based checks of actual exported DOCX and PDF.
These are local risk profiles, NOT vendor parser emulations or vendor scores.
"""
from io import BytesIO
import re
import zipfile
from collections import Counter
from docx import Document
from .documents import read_document

DOC_SOURCES = {
    'greenhouse': 'https://support.greenhouse.io/hc/en-us/articles/200989175-Unsuccessful-resume-parse',
    'workday': 'https://doc.workday.com/admin-guide/en-us/human-capital-management/recruiting/candidates/set-up-prospects-and-candidates/hdc1552497830785.html',
    'lever': 'https://help.lever.co/s/article/Understanding-Resume-Parsing',
}

def _check(key, title, outcome, reason, origin='local_test'):
    return {'id':key, 'title':title, 'status':outcome, 'detail':reason, 'evidence_type':origin}

def diagnose(docx_bytes: bytes, pdf_bytes: bytes | None = None) -> dict:
    """Analyse bytes of actually generated documents; never runs while typing."""
    d = Document(BytesIO(docx_bytes))
    text = read_document(docx_bytes, 'cv.docx')
    pdf_text = None
    if pdf_bytes is not None:
        try:
            pdf_text = read_document(pdf_bytes, 'cv.pdf')
        except Exception:
            pdf_text = ''
    tokens = lambda s: re.findall(r'\w+', s.casefold(), re.UNICODE)
    doc_tokens = Counter(tokens(text))
    pdf_tokens = Counter(tokens(pdf_text)) if pdf_text is not None else None
    if pdf_tokens is not None:
        missing = sum(max(0, count-pdf_tokens[token]) for token,count in doc_tokens.items())
        fidelity = round(100 * (1 - missing/max(1, sum(doc_tokens.values()))), 1)
    else:
        fidelity = None
    contact = bool(re.search(r'\b[^\s@]+@[^\s@]+\.[^\s@]+\b', text))
    sections = bool(re.search(r'(?i)doświadczenie|experience|employment',text))
    headings = [p.text.strip() for p in d.paragraphs if p.style and p.style.name.startswith('Heading')]
    structure = bool(headings) and sections
    tables = len(d.tables)
    images = len(d.inline_shapes)
    with zipfile.ZipFile(BytesIO(docx_bytes)) as z:
        xml = z.read('word/document.xml').decode('utf-8', errors='replace')
    text_boxes = xml.count('<w:txbxContent') + xml.count('<wps:txbx')
    header_contact = any('@' in p.text for s in d.sections for p in list(s.header.paragraphs)+list(s.footer.paragraphs))
    dates = bool(re.search(r'\b(?:19|20)\d\d(?:[-./](?:0?[1-9]|1[0-2]))?\b',text))
    tests = [
        _check('text', 'Tekst dokumentu daje się odczytać', 'pass' if len(text.strip()) >= 100 else 'warn', f'Odczytano {len(text.strip())} znaków z DOCX.'),
        _check('contact', 'Adres e-mail dostępny jako tekst', 'pass' if contact else 'warn', 'Adres jest rozpoznawalny.' if contact else 'Nie znaleziono poprawnego adresu e-mail.'),
        _check('sections', 'Sekcje mają czytelne nagłówki', 'pass' if structure else 'warn', f'Nagłówki Word: {len(headings)}; sekcja doświadczenia: {"tak" if sections else "nie"}.'),
        _check('layout', 'Brak tabel, grafik i pól tekstowych', 'pass' if not (tables or images or text_boxes) else 'warn', f'Tabele: {tables}, grafiki: {images}, pola tekstowe: {text_boxes}.'),
        _check('header', 'Dane kontaktowe poza nagłówkiem/stopką', 'warn' if header_contact else 'pass', 'Wykryto e-mail w nagłówku lub stopce.' if header_contact else 'Brak e-maila w nagłówku/stopce.'),
        _check('dates', 'Daty występują w odczytanym tekście', 'pass' if dates else 'warn', 'Wykryto rok lub datę.' if dates else 'Nie wykryto roku; sprawdź czy CV zawiera historię pracy.'),
        _check('size', 'Rozmiar DOCX poniżej 2,5 MB', 'pass' if len(docx_bytes)<=2_500_000 else 'fail', f'{len(docx_bytes)} bajtów; limit parsowania Greenhouse wynosi 2,5 MB.', 'vendor_documented'),
        _check('pdf', 'PDF można odczytać po eksporcie', 'unknown' if pdf_text is None else 'pass' if pdf_text.strip() else 'fail', 'Nie badano PDF (brak eksportu).' if pdf_text is None else f'Odczytano {len(pdf_text)} znaków z PDF.'),
        _check('roundtrip', 'Treść DOCX zachowana w PDF', 'unknown' if fidelity is None else 'pass' if fidelity >= 98 else 'warn' if fidelity >= 90 else 'fail', 'Nie badano PDF.' if fidelity is None else f'Pokrycie tokenów DOCX→PDF: {fidelity}% (próg lokalny 98%).'),
    ]
    lookup = {v['id']:v for v in tests}
    weights = {
        'greenhouse': {'text':20,'contact':10,'sections':15,'layout':20,'header':15,'size':10,'pdf':5,'roundtrip':5},
        'workday': {'text':20,'contact':15,'sections':15,'layout':10,'dates':20,'pdf':10,'roundtrip':10},
        'lever': {'text':25,'contact':15,'sections':10,'layout':10,'dates':10,'pdf':15,'roundtrip':15},
    }
    reports={}
    for vendor, items in weights.items():
        applicable=[(lookup[k],w) for k,w in items.items() if lookup[k]['status']!='unknown']
        maximum=sum(w for _,w in applicable)
        weighted=sum(w*(1 if c['status']=='pass' else 0.5 if c['status']=='warn' else 0) for c,w in applicable)
        reports[vendor]={
            'score':round(100*weighted/maximum) if maximum else None,
            'coverage':round(100*maximum/sum(items.values())),
            'label':'Autorski wskaźnik ryzyka technicznego (nie wynik ATS)',
            'rules':[dict(lookup[k],weight=w) for k,w in items.items()],
            'source':DOC_SOURCES[vendor],
        }
    return {
        'document_checks':tests,
        'profiles':reports,
        'pdf_checked':pdf_bytes is not None,
        'limitations':['Wyniki nie pochodzą z parserów Greenhouse, Workday ani Lever.',
                       'Wagi są heurystyką aplikacji i nie są ustalane przez producentów.',
                       'Brak PDF oznacza niepełną kontrolę; jego test jest wykonywany tylko na żądanie.',
                       'Wynik nie przewiduje odrzucenia, rankingu kandydata ani zgodności z konfiguracją pracodawcy.'],
    }
